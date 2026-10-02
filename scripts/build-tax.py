"""Build data/state-tax.json (ABS Taxation Revenue) and data/company-tax.json (ATO corporate tax transparency).

Needs openpyxl: pip install openpyxl
"""
import io, json, urllib.request
import openpyxl

ABS = "https://www.abs.gov.au/statistics/economy/government/taxation-revenue-australia/2024-25/55060DO001_202425.xlsx"
ATO = ("https://data.gov.au/data/dataset/c2524c87-cea4-4636-acac-599a82048a26/resource/"
       "15bf4eaa-7a14-48e9-8a3f-092df32a644e/download/2024-25-corporate-report-of-entity-tax-information.xlsx")
YEAR = "2024-25"


def workbook(url):
    return openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(url).read()), read_only=True)


# --- State and local government taxation, ABS 5506.0 tables 2-9 (state codes 1-8) ---
LINES = {
    "Taxes on employers payroll and labour force": "payroll",
    "Land taxes": "land",
    "Municipal rates": "rates",
    "Total taxes on gambling": "gambling",
    "Total taxes on insurance": "insurance",
    "Stamp duties on conveyances": "conveyance",
    "Total motor vehicle taxes": "motor",
}
wb = workbook(ABS)
states = {}
for code in range(1, 9):
    rows = list(wb[f"Table_{code + 1}"].iter_rows(values_only=True))
    years = [y for y in rows[4][1:] if y]
    data = {y: {} for y in years}
    for row in rows:
        label = (row[0] or "").strip()
        key = "total" if label.startswith("Total Taxation") else LINES.get(label)
        if key:
            for y, v in zip(years, row[1:]):
                data[y][key] = (v or 0) * 1e6
    states[str(code)] = data
json.dump({"source": "ABS Taxation Revenue, Australia, 2024-25 (5506.0), released 21 April 2026",
           "states": states}, open("data/state-tax.json", "w"), separators=(",", ":"))

# --- Company tax, ATO Report of Entity Tax Information 2024-25 ---
rows = list(workbook(ATO)["Income tax details"].iter_rows(values_only=True))[1:]
cos = [{"name": r[0], "abn": str(r[1]), "income": r[2] or 0, "taxable": r[3] or 0, "tax": r[4] or 0}
       for r in rows if r[5] == YEAR]
pick = lambda items, key: [dict(c) for c in sorted(items, key=key)[:15]]
json.dump({
    "source": f"ATO Corporate tax transparency, Report of Entity Tax Information {YEAR}",
    "year": YEAR,
    "count": len(cos),
    "totalTax": sum(c["tax"] for c in cos),
    "totalIncome": sum(c["income"] for c in cos),
    "noTaxCount": sum(1 for c in cos if not c["tax"]),
    "mostTax": pick(cos, lambda c: -c["tax"]),
    "noTaxBiggest": pick([c for c in cos if not c["tax"]], lambda c: -c["income"]),
    # Taxable income of $100m+ paying the smallest share of it in tax.
    "lowRate": pick([c for c in cos if c["taxable"] >= 1e8], lambda c: (c["tax"] / c["taxable"], -c["taxable"])),
}, open("data/company-tax.json", "w"), separators=(",", ":"))

assert len(states) == 8 and all(YEAR in s for s in states.values())
assert cos and all(c["tax"] <= c["taxable"] for c in cos if c["taxable"])
